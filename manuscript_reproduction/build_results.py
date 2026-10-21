"""Build the manuscript section from verified comparison.csv and audit results."""
from pathlib import Path
import json, math, re
import numpy as np
import pandas as pd
ROOT=Path(__file__).resolve().parent; OUT=ROOT/'revision_results'
d=pd.read_csv(OUT/'comparison.csv'); month=pd.read_csv(OUT/'monthly.csv')
audit=json.loads((OUT/'audit_revision.json').read_text()); assert audit['status']=='passed'
def v(case,metric='Total',col='iso_kWh'):
    return float(d[(d['case']==case)&(d.metric==metric)][col].iloc[0])
def f(x): return f'{x:.2f}'
def change(x): return 'increased' if x>0 else 'decreased'
blocks=[]
def title(s): blocks.append(('h1',s))
def sub(s): blocks.append(('h2',s))
def para(s): blocks.append(('p',s))
def table(caption,headers,rows): blocks.append(('table',(caption,headers,rows)))
def row(c,label): return [label,f(v(c,'Heating')),f(v(c,'Cooling')),f(v(c)),f(v(c,col='bias_pct')),f(v(c,col='RMSE_W'))]

title('4 Results')
para('The three-stage comparison separates implementation repairs, changes to solar transmission and exterior convection, and changes to Australian leakage and construction inputs. Twelve annual pyBuildingEnergy simulations and seven EnergyPlus 24.1.0 references were completed. All reported energies are sensible heating and cooling needs, rather than electricity consumption. Negative bias denotes underprediction relative to the corresponding EnergyPlus result.')
para('Stages 1 and 2 use the diagnostic zero-capacity construction and the historical low-leakage assumption of q₅₀ = 4 m³/(h m²). Stage 3 changes permeability to the provisional 2006 default of 11 and introduces material-derived construction properties, separately and together. Weather-dependent infiltration is prescribed identically to both engines using a 20 °C reference temperature for the stack driver. EnergyPlus algorithms remain unchanged; a new reference is calculated when physical scenario inputs change.')

sub('4.1 Implementation repairs')
para(f'The original implementation predicted {f(v("original"))} kWh of annual sensible demand against {f(v("original",col="ep_kWh"))} kWh in the common EnergyPlus reference. After the implementation repairs, demand was {f(v("repaired"))} kWh. Annual total bias changed from +{f(v("original",col="bias_pct"))}% to {f(v("repaired",col="bias_pct"))}%, while hourly total-load RMSE fell from {f(v("original",col="RMSE_W"))} to {f(v("repaired",col="RMSE_W"))} W (Table 1). This reduction occurred before either optional physical modification was introduced.')
table('Table 1. Implementation comparison against one EnergyPlus reference. Energy is kWh and RMSE is W. The infiltration-off case deliberately omits a pathway included in the reference.',
      ['Configuration','Heating','Cooling','Total','Bias %','RMSE'],
      [row('original','Original'),row('without_infiltration','Repaired without infiltration'),row('repaired','Repaired'),
       ['EnergyPlus',f(v('repaired','Heating','ep_kWh')),f(v('repaired','Cooling','ep_kWh')),f(v('repaired',col='ep_kWh')),'—','—']])
g0=audit['cases']['original']['internal_gains_kWh'];g1=audit['cases']['repaired']['internal_gains_kWh']
para(f'The diagnostic records confirmed that annual internal gains decreased from {f(g0)} to {f(g1)} kWh after repeated additions were removed. The original gain total was {g0/g1:.2f} times the repaired value. The original output also recorded {f(audit["cases"]["original"]["ground_flow_absolute_kWh"])} kWh of absolute ground heat flow for a dwelling with no ground contact; this quantity was zero after repair. These records identify errors in the represented gains and boundaries. They do not measure an equal reduction in conditioning demand, because the timing of gains, temperatures and thermostat operation affects the resulting heating and cooling.')
dh=v('repaired','Heating')-v('without_infiltration','Heating');dc=v('repaired','Cooling')-v('without_infiltration','Cooling')
para(f'In the controlled infiltration ablation, enabling the prescribed leakage flow increased heating by {f(dh)} kWh and cooling by {f(dc)} kWh, giving {f(dh+dc)} kWh of additional sensible demand. The repaired airflow calculation used only the 13.5 m² outdoor-exposed envelope, and its hourly conductance matched the independently calculated ventilation-plus-infiltration series. This verifies inclusion of the airflow pathway without establishing the real apartment’s permeability.')
para('The remaining repairs were assessed as a package. Focused checks supported internal-gain accounting, conditioned adjacency, ground-contact routing and heat-transfer calculations, but separate annual-demand effects were not assigned to every repair. The ground-season correction was inactive after ground-contact area was set to zero. The repaired model retained a substantial annual underprediction, so removal of the implementation faults did not establish equivalence with EnergyPlus.')

sub('4.2 Angular transmittance and exterior convection')
table('Table 2. Four physical-method combinations after implementation repairs. All use the same geometry, gains, airflow and EnergyPlus reference. Energy is kWh and RMSE is W.',
      ['Configuration','Heating','Cooling','Total','Bias %','RMSE'],
      [row('repaired','Fixed methods'),row('angular_only','Angular only'),row('convection_only','Convection only'),row('both','Both')])
angular=v('angular_only','Cooling')-v('repaired','Cooling');conv=v('convection_only','Cooling')-v('repaired','Cooling')
para(f'Angular transmittance alone reduced cooling by {f(-angular)} kWh, from {f(v("repaired","Cooling"))} to {f(v("angular_only","Cooling"))} kWh. Its heating effect was {v("angular_only","Heating")-v("repaired","Heating"):+.2f} kWh. Dynamic convection alone increased cooling by {f(conv)} kWh and changed heating by {v("convection_only","Heating")-v("repaired","Heating"):+.2f} kWh. The two cooling effects therefore acted in opposite directions.')
para(f'With both modifications enabled, cooling was {f(v("both","Cooling"))} kWh and total demand was {f(v("both"))} kWh. Absolute annual total bias increased from {abs(v("repaired",col="bias_pct")):.2f}% to {abs(v("both",col="bias_pct")):.2f}%. Dynamic convection alone produced the smallest absolute annual total bias of these four configurations, at {abs(v("convection_only",col="bias_pct")):.2f}%. The more detailed combined formulation therefore did not provide the closest annual result.')
para(f'Hourly agreement gave a different ranking. Both modifications reduced total-load RMSE from {f(v("repaired",col="RMSE_W"))} to {f(v("both",col="RMSE_W"))} W. Heating RMSE changed from {f(v("repaired","Heating","RMSE_W"))} to {f(v("both","Heating","RMSE_W"))} W, while cooling RMSE changed from {f(v("repaired","Cooling","RMSE_W"))} to {f(v("both","Cooling","RMSE_W"))} W. Annual bias and hourly error should consequently be interpreted separately.')
win=month[(month['case']=='repaired')&(month.metric=='Heating')&month.month.isin([6,7,8])]
short=float((win.ep_kWh-win.iso_kWh).sum());whole=v('repaired','Heating','ep_kWh')-v('repaired','Heating')
para(f'The monthly comparison located much of the heating difference in winter. June–August heating was {f(win.iso_kWh.sum())} kWh in the repaired model and {f(win.ep_kWh.sum())} kWh in EnergyPlus. The winter shortfall of {f(short)} kWh accounted for {100*short/whole:.1f}% of the annual heating difference. The optional solar and convection changes therefore did not remove the underlying heating discrepancy.')

sub('4.3 Australian leakage and construction inputs')
names=['both','au_leakage_only','au_construction_only','au_combined']
labels=['Corrected reference','Australian leakage only','Australian construction only','Combined Australian inputs']
table('Table 3. Four Australian input configurations. Both Stage 2 modifications are retained. Energy is kWh. Each row has a matched EnergyPlus reference; the construction change replaces diagnostic zero capacity with material-derived resistance and capacity.',
      ['Configuration','ISO heat','ISO cool','ISO total','EP total','Bias %'],
      [[l,f(v(c,'Heating')),f(v(c,'Cooling')),f(v(c)),f(v(c,col='ep_kWh')),f(v(c,col='bias_pct'))] for c,l in zip(names,labels)])
delta_leak=v('au_leakage_only')-v('both')
para(f'Increasing q₅₀ from 4 to 11 with diagnostic construction retained increased pyBuildingEnergy demand from {f(v("both"))} to {f(v("au_leakage_only"))} kWh, a rise of {f(delta_leak)} kWh. The corresponding EnergyPlus total increased from {f(v("both",col="ep_kWh"))} to {f(v("au_leakage_only",col="ep_kWh"))} kWh. Leakage changed demand in both engines, while a discrepancy remained. The age-band input therefore affects the predicted load but does not by itself validate the leakage assumption or resolve the model difference.')
dcst=v('au_construction_only')-v('both')
para(f'At q₅₀ = 4, replacing the diagnostic construction with the illustrative material-derived construction {change(dcst)} total demand by {f(abs(dcst))} kWh, to {f(v("au_construction_only"))} kWh. The corresponding EnergyPlus total was {f(v("au_construction_only",col="ep_kWh"))} kWh. This comparison changes both resistance and heat storage. It cannot isolate the influence of thermal capacity alone, and the reference case is a diagnostic construction rather than a measured alternative dwelling.')
para(f'The combined Australian input configuration predicted {f(v("au_combined","Heating"))} kWh of heating and {f(v("au_combined","Cooling"))} kWh of cooling, giving {f(v("au_combined"))} kWh in total against {f(v("au_combined",col="ep_kWh"))} kWh in EnergyPlus. Total annual bias was {v("au_combined",col="bias_pct"):+.2f}%, with hourly total-load RMSE of {f(v("au_combined",col="RMSE_W"))} W. Its absolute annual bias was {"smaller" if abs(v("au_combined",col="bias_pct"))<abs(v("both",col="bias_pct")) else "larger"} than that of the corrected reference configuration. Because the physical scenarios differ, this change in bias is a scenario-dependent comparison, not an isolated improvement to the numerical method.')
interaction=v('au_combined')-v('au_leakage_only')-v('au_construction_only')+v('both')
para(f'The combined demand change was not exactly the sum of the two separate changes. The interaction, calculated as combined minus leakage-only minus construction-only plus the reference result, was {interaction:+.2f} kWh in pyBuildingEnergy. This difference reflects the dependence of leakage heat transfer and thermostat operation on the changed construction response.')
para('The material-derived configuration retained the existing ISO capacity-class method. It did not use the earlier experimental layer-allocation extension. Consequently, the new results evaluate the effects of changed construction properties; they do not demonstrate improved representation of layer order. The material assemblies and neighbouring temperatures remain assumed case inputs.')

sub('4.4 Leakage sensitivity and airflow timing')
sens=['au_leakage_only','q14_sensitivity','q11_constant'];slabels=['q₅₀ 11 weather flow','q₅₀ 14 weather flow','q₅₀ 11 constant flow']
table('Table 4. Leakage sensitivity with diagnostic construction and both physical modifications. Energy is kWh. Weather-flow cases share the same prescribed hourly multiplier.',
      ['Scenario','ISO heat','ISO cool','EP heat','EP cool'],
      [[l,f(v(c,'Heating')),f(v(c,'Cooling')),f(v(c,'Heating','ep_kWh')),f(v(c,'Cooling','ep_kWh'))] for c,l in zip(sens,slabels)])
wd=v('au_leakage_only','Heating')-v('q11_constant','Heating')
mult=pd.read_csv(OUT/'hourly/diagnostic_q11_energyplus.csv').airflow_multiplier.mean()
para(f'At q₅₀ = 11, weather-dependent infiltration {change(wd)} heating by {f(abs(wd))} kWh relative to constant infiltration in pyBuildingEnergy. The annual mean multiplier was {mult:.3f}. Thus, annual mean airflow alone did not determine the heating response; the coincidence of leakage with outdoor temperature and heating operation also mattered. Raising q₅₀ from 11 to 14 changed total sensible demand by {v("q14_sensitivity")-v("au_leakage_only"):+.2f} kWh.')
para('These leakage effects are conditional on the small outdoor-exposed envelope and continuous designed ventilation of 0.040 m³/s. They should not be generalised to detached dwellings or more exposed apartments. No measured airflow data were used to calibrate either permeability or the weather multiplier.')

sub('4.5 Numerical verification and interpretation')
maxres=max(x.get('max_balance_residual_W',0) for x in audit['cases'].values())
para(f'All eleven corrected annual runs produced 8760 finite hourly loads and passed the hourly energy-balance criterion of 10⁻⁶ W. The largest absolute residual was {maxres:.2e} W. The transmission inventory resolved every surface: seven computational entries in the diagnostic model, where the identical windows were grouped, and eight in the construction model, where surfaces were retained separately. Ground heat flow remained zero. Independent checks verified the airflow conductance against the prescribed hourly multiplier and confirmed continuous ventilation in all seven EnergyPlus references.')
para(f'A separate historical reproduction case, using q₅₀ = 14, constant infiltration and fixed physical methods, returned {f(v("archive_reproduction","Heating"))} kWh of heating and {f(v("archive_reproduction","Cooling"))} kWh of cooling against {f(v("archive_reproduction",col="ep_kWh"))} kWh in EnergyPlus. This reproduced the earlier results within their reported precision and was kept separate from the revised primary comparison. The existing test suite passed 240 checks with 17 skipped checks; nine additional checks passed for the revised age lookup, airflow and construction design.')
para('Numerical consistency does not establish accuracy against a real building. The comparisons retain assumed neighbouring temperatures, glazing, ventilation and construction properties. The humidity diagnostic is excluded because a validated transient indoor moisture model is not available. The results support the verified implementation repairs and quantify the regional input sensitivities, while retaining the remaining disagreement with EnergyPlus.')

# Figures generated from the same verified simulation tables.
figure_specs = {
 '4.1': ('fig1_comparison_design', 'Figure 1. Comparison design. Implementation repairs precede the physical-method comparison; both optional methods are retained for the Australian input scenarios. The leakage ablation is a diagnostic check, and individual annual effects are not isolated for every repair.'),
 '4.2': ('fig2_annual_demand', 'Figure 2. Annual sensible demand. (a) Original and repaired implementations against the common EnergyPlus reference. (b) Australian input scenarios with their matched references. Panel scales differ; construction changes both resistance and capacity.'),
 '4.3': ('fig3_physical_methods', 'Figure 3. Annual bias and hourly RMSE for the four physical-method combinations. Zero bias and lower RMSE indicate closer agreement. The combined methods reduce hourly error but do not minimise annual bias.'),
 '4.5': ('fig4_monthly_heating', 'Figure 4. Monthly heating demand for the repaired implementation and its EnergyPlus reference. Shading identifies June–August, which accounts for 69.6% of the annual heating shortfall.')
}
expanded=[]
for kind,x in blocks:
    if kind=='h2' and x[:3] in figure_specs:
        expanded.append(('figure',figure_specs[x[:3]]))
    expanded.append((kind,x))
blocks=expanded

# One structured source generates every written format.
(OUT/'results_blocks.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2))
md=[]
for kind,x in blocks:
    if kind in ['h1','h2']: md.append(('# ' if kind=='h1' else '## ')+x)
    elif kind=='p': md.append(x)
    elif kind=='figure': md.append('!['+x[1]+'](figures/'+x[0]+'.png)\n\n'+x[1])
    else:
        caption,heads,rows=x;md.append(caption+'\n\n'+'| '+' | '.join(heads)+' |\n'+'| '+' | '.join(['---']*len(heads))+' |\n'+'\n'.join('| '+' | '.join(r)+' |' for r in rows))
(ROOT/'AIB_Results_Recomputed.md').write_text('\n\n'.join(md)+'\n')

def esc(s):
    s=s.replace('₅₀','$_{50}$').replace('m³','m$^3$').replace('m²','m$^2$').replace('°C',r'\textdegree C').replace('10⁻⁶',r'$10^{-6}$')
    for a,b in [('–','--'),('—','---'),('’',"'"),('%',r'\%'),('&',r'\&')]: s=s.replace(a,b)
    return s
tex=['% Recomputed October results. Paste after Methodology. Requires booktabs, tabularx and graphicx. Keep the figures folder beside this file.']
for kind,x in blocks:
    if kind=='h1': tex.append(r'\section{Results}')
    elif kind=='h2': tex.append(r'\subsection{'+esc(x.split(' ',1)[1])+'}')
    elif kind=='p':tex.append(esc(x))
    elif kind=='figure':
        tex.extend([r'\begin{figure}[htbp]\centering',r'\includegraphics[width=\linewidth]{figures/'+x[0]+'.pdf}',r'\caption{'+esc(x[1].split('. ',1)[1])+'}',r'\label{fig:'+x[0]+'}',r'\end{figure}'])
    else:
        caption,heads,rows=x
        tex.extend([r'\begin{table}[htbp]\centering\small',r'\caption{'+esc(caption.split('. ',1)[1])+'}',r'\begin{tabularx}{\textwidth}{@{}X'+'r'*(len(heads)-1)+'@{}}',r'\toprule', ' & '.join(map(esc,heads))+r' \\',r'\midrule'])
        tex += [' & '.join(map(esc,row))+r' \\' for row in rows]
        tex.extend([r'\bottomrule',r'\end{tabularx}',r'\end{table}'])
(ROOT/'AIB_Results_Section.tex').write_text('\n\n'.join(tex)+'\n')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(.9)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1.03
doc.styles['Title'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if style.type==1:
        rpr=style.element.get_or_add_rPr()
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is not None:
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+key),None)
for kind,x in blocks:
    if kind=='h1':doc.add_paragraph(x,'Title')
    elif kind=='h2':doc.add_paragraph(x,'Heading 2')
    elif kind=='p':doc.add_paragraph(x)
    elif kind=='figure':
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'figures'/(x[0]+'.png')),width=Inches(6.45))
        cp=doc.add_paragraph(x[1]);cp.paragraph_format.space_after=Pt(9)
        for rr in cp.runs:rr.font.size=Pt(9);rr.italic=True
    else:
        caption,heads,rows=x
        cp=doc.add_paragraph(caption);cp.paragraph_format.keep_with_next=True
        for rr in cp.runs:rr.font.size=Pt(10);rr.italic=True
        t=doc.add_table(rows=1,cols=len(heads));t.autofit=False
        widths=[2.25]+[(6.47-2.25)/(len(heads)-1)]*(len(heads)-1)
        for j,width in enumerate(widths):t.columns[j].width=Inches(width)
        for j,h in enumerate(heads):t.rows[0].cells[j].text=h
        for row in rows:
            cells=t.add_row().cells
            for j,val in enumerate(row):cells[j].text=val
        for i,r in enumerate(t.rows):
            trpr=r._tr.get_or_add_trPr();nr=OxmlElement('w:cantSplit');trpr.append(nr)
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(r.cells):
                c.width=Inches(widths[j]);pr=c._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                pr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF4' if i==0 else 'FFFFFF');pr.append(shade)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    p.paragraph_format.keep_with_next=i<len(t.rows)-1
                    if j:p.alignment=2
                    for run in p.runs:run.font.size=Pt(9);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(ROOT/'AIB_Results_Recomputed.docx')
print('Created manuscript formats from verified outputs.')
