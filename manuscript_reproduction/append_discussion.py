from pathlib import Path
from docx import Document
from docx.shared import Pt
R=Path(__file__).resolve().parent
# Rebuild Results first to make repeated execution idempotent.
import runpy
runpy.run_path(str(R/'build_results.py'))
text=(R/'AIB_Discussion_Conclusion.md').read_text()
doc=Document(R/'AIB_Results_Recomputed.docx')

tex=[]
def esc(s):
 return s.replace('₅₀',r'$_{50}$').replace('m³',r'm$^3$').replace('m²',r'm$^2$').replace('°C',r'\textdegree C').replace('%',r'\%').replace('−','-').replace('–','--').replace('’',"'")
for block in text.strip().split('\n\n'):
 if block.startswith('# '):
  title=block[2:];heading=doc.add_paragraph(title,'Title');heading.paragraph_format.page_break_before=True;tex.append(r'\section{'+title.split(' ',1)[1]+'}')
 elif block.startswith('## '):
  title=block[3:];doc.add_paragraph(title,'Heading 2');tex.append(r'\subsection{'+esc(title.split(' ',1)[1])+'}')
 else:
  doc.add_paragraph(block);tex.append(esc(block))
doc.save(R/'AIB_Results_Recomputed.docx')
fragment='\n\n'.join(tex)+'\n'
(R/'AIB_Discussion_Conclusion.tex').write_text(fragment)
with (R/'AIB_Results_Section.tex').open('a') as f:f.write('\n'+fragment)
with (R/'AIB_Results_Recomputed.md').open('a') as f:f.write('\n'+text)
print('Appended Discussion and Conclusion to all manuscript formats.')
